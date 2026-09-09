from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.shop_excel import (
    SheetNotFoundError,
    efficiency_snapshot,
    find_column,
    list_sheets,
    read_sheet,
    update_cells,
)


def _save_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    wb = Workbook()
    default = wb.active
    names = list(sheets)
    default.title = names[0]
    for title, rows in sheets.items():
        ws = default if title == names[0] else wb.create_sheet(title)
        for row in rows:
            ws.append(list(row))
    wb.save(path)
    wb.close()
    return path


def test_list_sheets(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "shop.xlsx",
        {
            "Production": [["Machine"]],
            "Scrap": [["Part"]],
            "Downtime": [["Minutes"]],
        },
    )
    assert list_sheets(path) == ["Production", "Scrap", "Downtime"]


def test_read_named_sheet_and_missing_sheet_raises(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "named.xlsx",
        {
            "Production": [["Machine", "OEE"], ["CNC-1", 92]],
            "Scrap": [["Part", "Qty"], ["Piston", 3]],
        },
    )
    production = read_sheet(path, "Production")
    assert production["sheet"] == "Production"
    assert production["rows"][0]["Machine"] == "CNC-1"
    scrap = read_sheet(path, "Scrap")
    assert scrap["sheet"] == "Scrap"
    assert scrap["rows"][0]["Part"] == "Piston"
    with pytest.raises(SheetNotFoundError, match="DoesNotExist") as raised:
        read_sheet(path, "DoesNotExist")
    message = str(raised.value)
    assert "Production" in message and "Scrap" in message
    with pytest.raises(SheetNotFoundError):
        efficiency_snapshot(path, "Nope")
    with pytest.raises(SheetNotFoundError):
        update_cells(path, "Nope", [{"cell": "A1", "value": 1}])


def test_headers_and_row_dict_mapping(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "headers.xlsx",
        {
            "Production": [
                ["Machine", None, "OEE", ""],
                ["Press-1", "ignored-without-header", 91, "also-ignored"],
                [None, None, None, None],
                ["Press-2", "x", 88, "y"],
            ]
        },
    )
    payload = read_sheet(path, "Production")
    assert payload["headers"] == ["Machine", "", "OEE", ""]
    assert payload["keyed_headers"] == ["Machine", "OEE"]
    assert payload["rows"] == [
        {"Machine": "Press-1", "OEE": 91},
        {"Machine": "Press-2", "OEE": 88},
    ]
    assert payload["raw"][0][0] == "Machine"
    assert "ignored-without-header" not in payload["rows"][0]


def test_update_cells_round_trip_and_does_not_mutate_other_sheets(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "update.xlsx",
        {
            "Production": [["Machine", "OEE"], ["CNC-1", 92], ["CNC-2", 81]],
            "Scrap": [["Keep", "Me"], ["untouched", 7]],
        },
    )
    before_scrap = read_sheet(path, "Scrap")
    result = update_cells(
        path,
        "Production",
        [
            {"cell": "B2", "value": 93.5},
            {"row": 3, "col": 2, "value": 90},
            {"row": 2, "col": "A", "value": "CNC-1-rev"},
        ],
    )
    assert result["sheet"] == "Production"
    assert "B2" in result["updated"]
    after = read_sheet(path, "Production")
    assert after["rows"][0] == {"Machine": "CNC-1-rev", "OEE": 93.5}
    assert after["rows"][1]["OEE"] == 90
    after_scrap = read_sheet(path, "Scrap")
    assert after_scrap == before_scrap
    wb = load_workbook(path)
    try:
        assert wb["Scrap"]["A2"].value == "untouched"
        assert wb["Scrap"]["B2"].value == 7
        assert wb["Production"]["B2"].value == 93.5
    finally:
        wb.close()
    with pytest.raises(FileNotFoundError):
        update_cells(tmp_path / "missing.xlsx", "Production", [{"cell": "A1", "value": 1}])


def test_find_column_case_insensitive():
    headers = ["Machine", "OEE %", "Scrap"]
    assert find_column(headers, ["oee", "efficiency", "oee %"]) == "OEE %"
    assert find_column(headers, ["EFFICIENCY", "oee%"]) == "OEE %"
    assert find_column(["Efficiency"], ["oee", "efficiency"]) == "Efficiency"
    assert find_column(headers, ["downtime", "dt"]) is None
    assert find_column(["", None, "Qty"], ["qty"]) == "Qty"


def test_efficiency_snapshot_all_at_or_above_90_has_no_bottlenecks(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "oee-good.xlsx",
        {
            "Production": [
                ["Machine", "OEE"],
                ["CNC-1", 90],
                ["CNC-2", 94],
                ["Press-1", 100],
            ]
        },
    )
    snap = efficiency_snapshot(path, "Production")
    assert snap["oee_percent"] == pytest.approx((90 + 94 + 100) / 3)
    assert snap["oee_percent"] >= 90
    assert snap["reason"] is None
    assert snap["bottlenecks"] == []
    assert snap["sample_size"] == 3
    assert snap["column"] == "OEE"


def test_efficiency_snapshot_mixed_oee_lists_bottlenecks(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "oee-mixed.xlsx",
        {
            "Production": [
                ["Machine", "Operation", "OEE"],
                ["CNC-1", "Turn", 95],
                ["Press-2", "Form", 70],
                ["Line-B", "Assemble", 88],
                ["CNC-2", "Mill", 92],
            ]
        },
    )
    snap = efficiency_snapshot(path, "Production")
    assert snap["oee_percent"] == pytest.approx((95 + 70 + 88 + 92) / 4)
    labels = [item["label"] for item in snap["bottlenecks"]]
    assert labels == ["Press-2", "Line-B"]
    by_label = {item["label"]: item for item in snap["bottlenecks"]}
    assert by_label["Press-2"]["oee_percent"] == pytest.approx(70)
    assert by_label["Line-B"]["oee_percent"] == pytest.approx(88)
    assert by_label["Press-2"]["row"] == 3


def test_efficiency_snapshot_no_oee_column_does_not_invent_numbers(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "no-oee.xlsx",
        {
            "Production": [
                ["Machine", "Parts", "Scrap"],
                ["CNC-1", 120, 2],
                ["CNC-2", 80, 1],
            ]
        },
    )
    snap = efficiency_snapshot(path, "Production")
    assert snap["oee_percent"] is None
    assert snap["oee_percent"] is not False
    assert snap["reason"] == "no efficiency column"
    assert snap["bottlenecks"] == []
    assert snap["sample_size"] == 0
    assert 0 not in (snap["oee_percent"],)
    assert snap["oee_percent"] != 0
    assert snap["oee_percent"] != 90


def test_ratio_zero_point_nine_two_versus_percent_ninety_two(tmp_path: Path):
    ratio_path = _save_workbook(
        tmp_path / "ratio.xlsx",
        {
            "Production": [
                ["Machine", "Efficiency"],
                ["A", 0.92],
                ["B", 0.88],
                ["C", 0.95],
            ]
        },
    )
    percent_path = _save_workbook(
        tmp_path / "percent.xlsx",
        {
            "Production": [
                ["Machine", "Efficiency"],
                ["A", 92],
                ["B", 88],
                ["C", 95],
            ]
        },
    )
    ratio = efficiency_snapshot(ratio_path, "Production")
    percent = efficiency_snapshot(percent_path, "Production")
    expected = pytest.approx((92 + 88 + 95) / 3)
    assert ratio["oee_percent"] == expected
    assert percent["oee_percent"] == expected
    assert [item["label"] for item in ratio["bottlenecks"]] == ["B"]
    assert [item["label"] for item in percent["bottlenecks"]] == ["B"]
    assert ratio["bottlenecks"][0]["oee_percent"] == pytest.approx(88)
    assert percent["bottlenecks"][0]["oee_percent"] == pytest.approx(88)


def test_blank_and_non_numeric_cells_ignored_not_treated_as_zero(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "holes.xlsx",
        {
            "Production": [
                ["Machine", "OEE %"],
                ["CNC-1", 90],
                ["CNC-2", None],
                ["CNC-3", ""],
                ["CNC-4", "n/a"],
                ["CNC-5", "hold"],
                ["CNC-6", 80],
                ["CNC-7", "92%"],
            ]
        },
    )
    snap = efficiency_snapshot(path, "Production")
    assert snap["sample_size"] == 3
    assert snap["oee_percent"] == pytest.approx((90 + 80 + 92) / 3)
    assert snap["oee_percent"] != pytest.approx((90 + 0 + 0 + 0 + 0 + 80 + 92) / 7)
    labels = [item["label"] for item in snap["bottlenecks"]]
    assert labels == ["CNC-6"]
    assert "CNC-2" not in labels
    assert "CNC-4" not in labels

    empty_only = _save_workbook(
        tmp_path / "empty-oee.xlsx",
        {
            "Production": [
                ["Machine", "OEE"],
                ["CNC-1", None],
                ["CNC-2", "n/a"],
            ]
        },
    )
    empty_snap = efficiency_snapshot(empty_only, "Production")
    assert empty_snap["oee_percent"] is None
    assert empty_snap["reason"] == "no numeric efficiency values"
    assert empty_snap["column"] == "OEE"


def test_negative_oee_ignored_not_treated_as_zero(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "neg-oee.xlsx",
        {
            "Production": [
                ["Machine", "OEE"],
                ["CNC-1", 90],
                ["CNC-2", -5],
                ["CNC-3", 80],
            ]
        },
    )
    snap = efficiency_snapshot(path, "Production")
    assert snap["sample_size"] == 2
    assert snap["oee_percent"] == pytest.approx((90 + 80) / 2)
    labels = [item["label"] for item in snap["bottlenecks"]]
    assert labels == ["CNC-3"]
    assert "CNC-2" not in labels


def test_efficiency_alias_column_and_operation_label(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "alias.xlsx",
        {
            "Shift": [
                ["Operation", "efficiency"],
                ["Hone", 91],
                ["Wash", 0],
            ]
        },
    )
    snap = efficiency_snapshot(path, "Shift")
    assert snap["column"] == "efficiency"
    assert snap["oee_percent"] == pytest.approx(45.5)
    assert snap["bottlenecks"][0]["label"] == "Wash"
    assert snap["bottlenecks"][0]["oee_percent"] == pytest.approx(0)


def test_does_not_mutate_unexpected_sheets_on_read(tmp_path: Path):
    path = _save_workbook(
        tmp_path / "readonly.xlsx",
        {
            "Production": [["Machine", "OEE"], ["A", 91]],
            "Notes": [["Secret", "Value"], ["do-not-touch", 42]],
        },
    )
    read_sheet(path, "Production")
    efficiency_snapshot(path, "Production")
    list_sheets(path)
    wb = load_workbook(path)
    try:
        assert wb.sheetnames == ["Production", "Notes"]
        assert wb["Notes"]["A2"].value == "do-not-touch"
        assert wb["Notes"]["B2"].value == 42
        assert wb["Production"]["B2"].value == 91
    finally:
        wb.close()
