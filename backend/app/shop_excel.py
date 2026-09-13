"""Read and update production workbooks by named sheet.

Pure file-based helpers. Never invent OEE or other shop numbers: missing
efficiency columns and blank/non-numeric cells stay missing.
"""

from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils.exceptions import InvalidFileException

DEFAULT_OEE_ALIASES: tuple[str, ...] = ("oee", "efficiency", "oee %")
_LABEL_ALIASES: tuple[str, ...] = (
    "machine",
    "machine name",
    "operation",
    "op",
    "name",
    "line",
    "station",
    "equipment",
    "asset",
    "work center",
    "workcentre",
    "work cell",
)
_SPACE_RE = re.compile(r"\s+")


class SheetNotFoundError(ValueError):
    """Named sheet does not exist; the active sheet is never used as a fallback."""


def _as_path(path: str | Path) -> Path:
    return Path(path)


def _load(path: str | Path, *, data_only: bool = False, read_only: bool = False):
    target = _as_path(path)
    if not target.is_file():
        raise FileNotFoundError(f"Workbook not found: {target}")
    try:
        return load_workbook(target, data_only=data_only, read_only=read_only)
    except (InvalidFileException, KeyError, OSError) as exc:
        raise ValueError(f"Cannot open workbook: {target}") from exc


def _require_sheet(wb, sheet_name: str):
    names = list(wb.sheetnames)
    if sheet_name not in names:
        available = ", ".join(names) if names else "(none)"
        raise SheetNotFoundError(
            f"Sheet {sheet_name!r} not found. Available: {available}"
        )
    return wb[sheet_name]


def _header_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    text = _header_text(value).lower().replace("%", " ")
    text = _SPACE_RE.sub(" ", text).strip()
    return text


def _cell_value(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    return value


def _row_is_empty(values: list[Any]) -> bool:
    for value in values:
        if value is None:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        return False
    return True


def _record_from_row(headers: list[str], values: list[Any]) -> dict[str, Any]:
    record: dict[str, Any] = {}
    for index, header in enumerate(headers):
        if not header or header in record:
            continue
        record[header] = values[index] if index < len(values) else None
    return record


def list_sheets(path: str | Path) -> list[str]:
    wb = _load(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_sheet(path: str | Path, sheet_name: str) -> dict[str, Any]:
    """Read one named sheet. Raises SheetNotFoundError if the name is missing.

    Uses cached cell values (data_only) so formula results are read when Excel
    has calculated them. Uncalculated formulas stay missing — never invented.
    """
    wb = _load(path, data_only=True, read_only=True)
    try:
        ws = _require_sheet(wb, sheet_name)
        raw: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            raw.append([_cell_value(cell) for cell in row])
        while raw and _row_is_empty(raw[-1]):
            raw.pop()
        headers = [_header_text(cell) for cell in raw[0]] if raw else []
        keyed_headers = [name for name in headers if name]
        rows: list[dict[str, Any]] = []
        for values in raw[1:]:
            if _row_is_empty(values):
                continue
            rows.append(_record_from_row(headers, values))
        return {
            "sheet": sheet_name,
            "headers": headers,
            "rows": rows,
            "raw": raw,
            "keyed_headers": keyed_headers,
        }
    finally:
        wb.close()


def find_column(headers: list[Any], aliases: list[str] | tuple[str, ...]) -> str | None:
    """Return the first header that case-insensitively matches an alias, else None."""
    index: dict[str, str] = {}
    for header in headers:
        name = _header_text(header)
        key = _normalize_header(name)
        if not name or not key or key in index:
            continue
        index[key] = name
    for alias in aliases:
        key = _normalize_header(alias)
        if key and key in index:
            return index[key]
    return None


def _parse_update(item: dict[str, Any]) -> tuple[int, int, Any]:
    if not isinstance(item, dict):
        raise TypeError("Each update must be a dict")
    if "value" not in item:
        raise ValueError("Each update requires a 'value'")
    value = item["value"]
    cell = item.get("cell")
    if cell:
        col_letter, row = coordinate_from_string(str(cell).strip())
        return int(row), column_index_from_string(col_letter), value
    if "row" in item and "col" in item:
        row = int(item["row"])
        col = item["col"]
        if isinstance(col, str) and col.strip() and not col.strip().isdigit():
            col_index = column_index_from_string(col.strip())
        else:
            col_index = int(col)
        if row < 1 or col_index < 1:
            raise ValueError("row and col are 1-based Excel coordinates")
        return row, col_index, value
    raise ValueError("Each update requires 'cell' or both 'row' and 'col'")


def update_cells(
    path: str | Path,
    sheet_name: str,
    updates: list[dict[str, Any]],
) -> dict[str, Any]:
    """Write values on an existing named sheet and save in place.

    Does not create a workbook or a sheet. Missing paths raise FileNotFoundError.
    """
    target = _as_path(path)
    wb = _load(target)
    try:
        ws = _require_sheet(wb, sheet_name)
        written: list[str] = []
        for item in updates:
            row, col, value = _parse_update(item)
            ws.cell(row=row, column=col, value=value)
            written.append(f"{get_column_letter(col)}{row}")
        wb.save(target)
        return {
            "sheet": sheet_name,
            "path": str(target),
            "updated": written,
        }
    finally:
        wb.close()


def _to_number(value: Any) -> float | None:
    """Parse a numeric cell. Blanks and non-numerics are skipped, never coerced to 0."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        if not math.isfinite(number) or number < 0:
            return None
        return number
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None
        if text.endswith("%"):
            text = text[:-1].strip()
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
        if not math.isfinite(number) or number < 0:
            return None
        return number
    return None


def _to_percent_scale(values: list[float]) -> list[float]:
    if values and all(value <= 1 for value in values):
        return [value * 100.0 for value in values]
    return list(values)


def _label_for_row(row: dict[str, Any], label_header: str | None, excel_row: int) -> str:
    if label_header:
        raw = row.get(label_header)
        if raw is not None and not (isinstance(raw, str) and not raw.strip()):
            return str(raw).strip()
    return f"row {excel_row}"


def efficiency_from_payload(
    payload: dict[str, Any],
    oee_aliases: tuple[str, ...] | list[str] = DEFAULT_OEE_ALIASES,
) -> dict[str, Any]:
    """Mean OEE from a table payload. Never invents a number."""
    headers = payload.get("headers") or []
    column = find_column(headers, list(oee_aliases))
    if column is None:
        return {
            "oee_percent": None,
            "reason": "no efficiency column",
            "bottlenecks": [],
            "column": None,
            "sample_size": 0,
        }

    samples: list[tuple[int, dict[str, Any], float]] = []
    # Header is Excel row 1; data rows start at 2. Empty data rows were skipped
    # by read_sheet, so recover the Excel row from raw when possible.
    raw: list[list[Any]] = payload.get("raw") or []
    header_index = {name: idx for idx, name in enumerate(headers) if name}
    col_idx = header_index.get(column)
    excel_row = 1
    for values in raw[1:]:
        excel_row += 1
        if _row_is_empty(values):
            continue
        cell = values[col_idx] if col_idx is not None and col_idx < len(values) else None
        number = _to_number(cell)
        if number is None:
            continue
        samples.append((excel_row, _record_from_row(headers, values), number))

    if not samples:
        return {
            "oee_percent": None,
            "reason": "no numeric efficiency values",
            "bottlenecks": [],
            "column": column,
            "sample_size": 0,
        }

    percents = _to_percent_scale([sample[2] for sample in samples])
    mean = sum(percents) / len(percents)
    label_header = find_column(headers, _LABEL_ALIASES)
    bottlenecks: list[dict[str, Any]] = []
    for (excel_row, record, _raw_number), percent in zip(samples, percents):
        if percent < 90:
            bottlenecks.append(
                {
                    "label": _label_for_row(record, label_header, excel_row),
                    "oee_percent": percent,
                    "row": excel_row,
                }
            )
    return {
        "oee_percent": mean,
        "reason": None,
        "bottlenecks": bottlenecks,
        "column": column,
        "sample_size": len(samples),
    }


def efficiency_snapshot(
    path: str | Path,
    sheet_name: str,
    oee_aliases: tuple[str, ...] | list[str] = DEFAULT_OEE_ALIASES,
) -> dict[str, Any]:
    """Mean OEE from a real column, or None if that column is absent.

    Never fills oee_percent with a made-up number. Blank and non-numeric cells
    are ignored (not treated as 0). Values are treated as a 0–1 ratio only when
    every present number is <= 1; otherwise they are already percent.
    """
    return efficiency_from_payload(read_sheet(path, sheet_name), oee_aliases)
