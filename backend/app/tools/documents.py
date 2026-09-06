from __future__ import annotations

import csv
import io
import uuid
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from ..config import settings
from ..db import add_artifact, safe_export_path


def _unique_name(stem: str, ext: str) -> str:
    clean = "".join(ch if ch.isalnum() or ch in "-_ " else "-" for ch in stem).strip() or "jarvis"
    return f"{clean}-{uuid.uuid4().hex[:6]}.{ext}"


def create_spreadsheet(
    title: str,
    columns: list[str],
    rows: list[list[Any]],
    chart: bool = False,
) -> dict[str, Any]:
    name = _unique_name(title or "workbook", "xlsx")
    path = safe_export_path(name)
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Sheet")[:31]
    header_font = Font(bold=True, color="0B1F33")
    ws.append(columns)
    for cell in ws[1]:
        cell.font = header_font
    for row in rows:
        ws.append(list(row))
    if chart and columns and rows:
        chart_obj = BarChart()
        chart_obj.title = title
        data = Reference(ws, min_col=2, min_row=1, max_col=len(columns), max_row=len(rows) + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=len(rows) + 1)
        chart_obj.add_data(data, titles_from_data=True)
        chart_obj.set_categories(cats)
        ws.add_chart(chart_obj, "E2")
    wb.save(path)
    return add_artifact(uuid.uuid4().hex[:12], "xlsx", name, str(path))


def create_document(title: str, body: str, bullets: list[str] | None = None) -> dict[str, Any]:
    name = _unique_name(title or "document", "docx")
    path = safe_export_path(name)
    doc = Document()
    doc.add_heading(title or "Jarvis note", level=1)
    for paragraph in (body or "").split("\n\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())
    if bullets:
        for item in bullets:
            doc.add_paragraph(item, style="List Bullet")
    doc.save(path)
    return add_artifact(uuid.uuid4().hex[:12], "docx", name, str(path))


def create_pdf(title: str, body: str) -> dict[str, Any]:
    name = _unique_name(title or "brief", "pdf")
    path = safe_export_path(name)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(title or "Jarvis brief", styles["Title"]),
        Spacer(1, 12),
    ]
    for paragraph in (body or "").split("\n"):
        story.append(Paragraph(paragraph or "&nbsp;", styles["BodyText"]))
        story.append(Spacer(1, 8))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)
    return add_artifact(uuid.uuid4().hex[:12], "pdf", name, str(path))


def spreadsheet_from_csv(title: str, csv_text: str) -> dict[str, Any]:
    reader = csv.reader(io.StringIO(csv_text))
    table = list(reader)
    if not table:
        raise ValueError("CSV is empty")
    return create_spreadsheet(title, table[0], table[1:])


def read_export_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt" or suffix == ".md":
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".csv":
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            pages.append((page.extract_text() or "").strip())
        return "\n\n".join(part for part in pages if part)
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        lines = []
        for row in ws.iter_rows(values_only=True):
            lines.append(", ".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(lines)
    if suffix == ".docx":
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    return path.read_text(encoding="utf-8", errors="replace")


def workspace_root() -> Path:
    return settings.exports_dir
